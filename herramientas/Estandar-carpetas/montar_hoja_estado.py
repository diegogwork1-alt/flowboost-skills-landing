#!/usr/bin/env python3
"""Monta la hoja de ESTADO DE CUENTA de un cliente: en qué punto del proceso está.

Es la foto del proceso, no de los números (los números están en la hoja de reportes).
Todas las etapas del funnel vienen **ya escritas**, con su carril y su dueño: no hay que
redactar nada. El estado se elige de un **desplegable** (Pendiente · En curso · Hecho ·
Bloqueado · No aplica) y la hoja se pinta sola según lo elegido.

Arriba, un semáforo que se calcula solo: cuántas etapas hechas, si el carril de lanzamiento
está listo y qué está bloqueado.

Como con la hoja de reportes: se compone con openpyxl y se sube a Drive **con conversión**,
así que en Drive queda un Sheet nativo. La API de Sheets no está disponible con el token de
rclone (403: el proyecto no la tiene habilitada).

Uso:
  python3 montar_hoja_estado.py --cliente Cliente 02 --carpeta-id <ID>
"""
import argparse, os, sys, tempfile, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from subir_a_drive import api, _multipart

HOJA = "application/vnd.google-apps.spreadsheet"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MORADO, ROSA, NEGRO, GRIS = "7C5CFF", "FF2E7E", "0B0B0F", "8A8A94"
SUAVE, BLANCO = "FAF9FC", "FFFFFF"
VERDE, VERDEBG = "16794A", "EAF3EE"
AMBAR, AMBARBG = "9A7A12", "FFF6E0"
ROJO, ROJOBG = "C0392B", "FDEBE9"
AZULBG = "EEF2FF"

ESTADOS = ["Pendiente", "En curso", "Hecho", "Bloqueado", "No aplica"]

# (carril, etapa, dueño, qué hay que ver para darla por hecha)
ETAPAS = [
 ("Arranque", "Carpeta del cliente en Drive", "Operaciones", "Existe i_<C>/c_<C> con las numeradas"),
 ("Arranque", "Brief desde el onboarding", "Agente", "Brief_<C>.pdf en 0. Onboarding"),
 ("Arranque", "Mensaje de accesos enviado", "Agente", "Copia en Mensajes/ con fecha"),
 ("Arranque", "Accesos recibidos (Meta, DNS)", "Cliente", "Cuenta de Meta accesible en solo lectura"),
 ("Arranque", "Alta en reportes (día 2)", "Agente", "Su hoja de reportes creada y en la lista"),
 ("Arranque", "Radar de competencia", "Agente", "Documento en Drive"),

 ("A · Lanzamiento", "Estáticos de lanzamiento (mín. 6)", "Agente", "6 piezas aprobadas, 1:1 y 9:16"),
 ("A · Lanzamiento", "Copy de anuncios", "Agente", "Copy por estático, con variantes"),
 ("A · Lanzamiento", "Landing SIN VSL (copy + imagen)", "Dirección", "Maqueta aprobada"),
 ("A · Lanzamiento", "Landing publicada en el subdominio", "Agente", "Responde 200 en su URL"),
 ("A · Lanzamiento", "Formulario", "Dirección", "Se envía y llega el aviso"),
 ("A · Lanzamiento", "Thank you page", "Dirección", "Existe y carga tras enviar"),
 ("A · Lanzamiento", "Píxel instalado", "Agente", "Se ve en la landing o en GTM"),
 ("A · Lanzamiento", "Evento Lead comprobado", "Agente", "Salta en el Test de Eventos"),
 ("A · Lanzamiento", "Lead de prueba de punta a punta", "Agente", "Llega el aviso y queda registrado"),
 ("A · Lanzamiento", "Campaña armada (en pausa)", "Agente", "Estructura montada, sin activar"),
 ("A · Lanzamiento", "CAMPAÑA ACTIVA", "Dirección", "OK explícito de Dirección"),

 ("B · Vídeo", "5 guiones EGC escritos y auditados", "Agente", "PDF en 2. Ads/EGC/Guiones"),
 ("B · Vídeo", "Guion del VSL escrito y auditado", "Agente", "PDF en 2. Ads/EGC/Guiones"),
 ("B · Vídeo", "Enviado al cliente (EGC + VSL + guía)", "Agente", "Un solo mensaje, copia en Mensajes/"),
 ("B · Vídeo", "Vídeos crudos recibidos", "Cliente", "Archivos en 2. Ads/Vídeos crudos"),
 ("B · Vídeo", "Anuncios 9:16 editados", "Agente", "En 2. Ads/EGC/Editados"),
 ("B · Vídeo", "VSL editado", "Agente", "Vídeo final aprobado"),
 ("B · Vídeo", "VSL añadido a la landing", "Agente", "Se ve en la landing publicada"),
 ("B · Vídeo", "Vídeos metidos en la cuenta", "Agente", "Sustituyendo a lo que se mató"),

 ("C · Operación", "Revisión semanal de cuenta", "Agente", "Tabla matar/mantener/escalar"),
 ("C · Operación", "Informe de landing (semanal)", "Agente", "Fila en el histórico"),
 ("C · Operación", "Retro mensual de creatividades", "Agente", "Días 1-5 del mes"),
 ("C · Operación", "Reportes automáticos funcionando", "Agente",
  "El agente comprueba el viernes que salió el correo y lo marca"),
]

COLS = [("Carril", 16), ("Etapa", 40), ("Estado", 15), ("Quién", 12),
        ("Fecha", 13), ("Qué hay que ver para darla por hecha", 44), ("Notas", 40)]


def F(ws, ref, v, *, tam=10, bold=False, color=NEGRO, fondo=None, hor=None, italic=False):
    c = ws[ref]; c.value = v
    c.font = Font(name="Inter", size=tam, bold=bold, color=color, italic=italic)
    if fondo: c.fill = PatternFill("solid", fgColor=fondo)
    c.alignment = Alignment(horizontal=hor, vertical="center")
    return c


# La cabecera de la tabla va en la fila 7 (debajo del título y el semáforo), y así se queda:
# es como se lee bien. n8n sabe encontrarla porque su nodo de Google Sheets se configura con
# `headerRow: 7`; no hace falta partir la hoja en dos por una limitación de una herramienta.
FILA_CABECERA = 7
FILA_PRIMERA = 8


def construir(cliente):
    wb = Workbook(); ws = wb.active; ws.title = "Estado"
    ws.sheet_properties.tabColor = MORADO
    ws.sheet_view.showGridLines = False
    for j, (_, an) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = an
    ncol = len(COLS); ult = get_column_letter(ncol)

    ws.merge_cells(f"A1:{ult}1"); ws.row_dimensions[1].height = 44
    F(ws, "A1", f"  {cliente.upper()}   ·   ESTADO DE CUENTA",
      tam=18, bold=True, color=BLANCO, fondo=NEGRO)
    for j in range(2, ncol+1):
        ws[f"{get_column_letter(j)}1"].fill = PatternFill("solid", fgColor=NEGRO)
    ws.merge_cells(f"A2:{ult}2")
    F(ws, "A2", "  En qué punto del proceso está. El estado se elige del desplegable: "
                "no hay que escribir nada. Lo del agente se marca solo.", tam=9, color=GRIS)

    primera, ultima = FILA_PRIMERA, FILA_PRIMERA + len(ETAPAS) - 1
    cont = lambda e: f'COUNTIF($C${primera}:$C${ultima},"{e}")'
    idx_A = [i for i, e in enumerate(ETAPAS) if e[0].startswith("A ·")]
    fA1 = primera + idx_A[0]
    idx_activa = next(i for i, e in enumerate(ETAPAS) if e[1] == "CAMPAÑA ACTIVA")
    fPrev = primera + idx_activa - 1
    nA = fPrev - fA1 + 1

    resumen = [
        ("HECHAS", f'=TEXT({cont("Hecho")},"0")&" de {len(ETAPAS)}"', VERDE, VERDEBG),
        ("EN CURSO", f'={cont("En curso")}', AMBAR, AMBARBG),
        ("BLOQUEADAS", f'={cont("Bloqueado")}', ROJO, ROJOBG),
        ("¿SE PUEDE LANZAR?",
         f'=IF(COUNTIF($C${fA1}:$C${fPrev},"Hecho")>={nA},"SÍ, todo listo",'
         f'"Faltan "&TEXT({nA}-COUNTIF($C${fA1}:$C${fPrev},"Hecho"),"0"))', MORADO, AZULBG),
    ]
    F(ws, "A4", "DE UN VISTAZO", tam=9, bold=True, color=GRIS)
    for j, (t, v, col, bg) in enumerate(resumen):
        c1 = get_column_letter(1 + j*2); c2 = get_column_letter(2 + j*2)
        ws.merge_cells(f"{c1}5:{c2}5"); ws.merge_cells(f"{c1}6:{c2}6")
        F(ws, f"{c1}5", t, tam=8, bold=True, color=GRIS, fondo=bg)
        F(ws, f"{c1}6", v, tam=14, bold=True, color=col, fondo=bg)
        for cc in (c1, c2):
            for r in (5, 6): ws[f"{cc}{r}"].fill = PatternFill("solid", fgColor=bg)
    ws.row_dimensions[6].height = 26

    for j, (t, _) in enumerate(COLS, start=1):
        F(ws, f"{get_column_letter(j)}{FILA_CABECERA}", t, tam=8, bold=True, color=BLANCO,
          fondo=NEGRO, hor="center" if j in (3, 4, 5) else "left")

    carril_previo = None
    for i, (carril, etapa, quien, prueba) in enumerate(ETAPAS):
        r = primera + i
        nuevo = carril != carril_previo; carril_previo = carril
        F(ws, f"A{r}", carril if nuevo else "", tam=8, bold=nuevo, color=MORADO if nuevo else GRIS)
        F(ws, f"B{r}", etapa, tam=9, bold=etapa.isupper())
        F(ws, f"C{r}", "Pendiente", tam=9, hor="center")
        F(ws, f"D{r}", quien, tam=9, color=GRIS, hor="center")
        F(ws, f"E{r}", "", tam=9, hor="center")
        F(ws, f"F{r}", prueba, tam=8, color=GRIS)
        F(ws, f"G{r}", "", tam=9)

    dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "Elegí uno de la lista"; dv.errorTitle = "Estado no válido"
    ws.add_data_validation(dv); dv.add(f"C{primera}:C{ultima}")

    rango = f"A{primera}:{ult}{ultima}"
    for valor, letra, bg in [("Hecho", VERDE, VERDEBG), ("En curso", AMBAR, AMBARBG),
                             ("Bloqueado", ROJO, ROJOBG), ("No aplica", GRIS, SUAVE)]:
        ws.conditional_formatting.add(rango, FormulaRule(
            formula=[f'$C{primera}="{valor}"'],
            fill=PatternFill("solid", fgColor=bg),
            font=Font(name="Inter", size=9, color=letra,
                      bold=(valor in ("Bloqueado", "Hecho")))))

    ws.freeze_panes = f"A{primera}"
    ws.auto_filter.ref = f"A{FILA_CABECERA}:{ult}{ultima}"
    r = ultima + 2
    F(ws, f"A{r}", "Lo que depende del AGENTE se marca solo al cerrar cada etapa. Lo de Dirección y "
                   "lo de Operaciones se marca a mano. Una etapa «Bloqueada» tiene que llevar en Notas "
                   "por qué y a quién espera.", tam=8, color=GRIS, italic=True)
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cliente", required=True)
    ap.add_argument("--carpeta-id", default="",
                    help="ID de la carpeta de reportes; si no se pasa, se busca con --cliente-raiz")
    ap.add_argument("--cliente-raiz", default="",
                    help="gdrive:i_X/c_X — encuentra la carpeta de reportes sola (5., 6. o 7.)")
    ap.add_argument("--nombre", default=None)
    ap.add_argument("--rehacer", action="store_true",
                    help="regenerar aunque exista. PISA lo que estuviera marcado.")
    a = ap.parse_args()

    # La carpeta de reportes no tiene número fijo: se busca por nombre.
    carpeta_id = a.carpeta_id
    if not carpeta_id:
        if not a.cliente_raiz:
            sys.exit("✗ Hace falta --carpeta-id o --cliente-raiz.")
        import os as _os
        sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
        from carpeta_reportes import buscar
        from subir_a_drive import folder_id
        carpeta = buscar(a.cliente_raiz)
        if not carpeta:
            sys.exit(f"✗ No hay carpeta de reportes en {a.cliente_raiz}. La crea Operaciones.")
        carpeta_id = folder_id(carpeta)
        print(f"✓ Carpeta de reportes: {carpeta}")

    nombre = a.nombre or f"Estado de cuenta — {a.cliente}"
    wb = construir(a.cliente)
    tmp = os.path.join(tempfile.mkdtemp(), "estado.xlsx"); wb.save(tmp)

    q = f"name = '{nombre}' and '{carpeta_id}' in parents and trashed = false"
    url = ("https://www.googleapis.com/drive/v3/files?"
           + urllib.parse.urlencode({"q": q, "fields": "files(id,name)",
                                     "supportsAllDrives": "true",
                                     "includeItemsFromAllDrives": "true"}))
    ya = api(url).get("files", [])
    if ya and not a.rehacer:
        # Regenerar PISA lo marcado. Por defecto NO se toca: se devuelve la que hay.
        print(f"= Ya existe, no se toca (tiene marcas dentro): {nombre}")
        print(f"  https://docs.google.com/spreadsheets/d/{ya[0]['id']}/edit")
        print(ya[0]["id"]); return
    if ya:
        sid = ya[0]["id"]
        cuerpo, cab = _multipart({"mimeType": HOJA}, tmp, XLSX)
        api(f"https://www.googleapis.com/upload/drive/v3/files/{sid}"
            f"?uploadType=multipart&supportsAllDrives=true",
            data=cuerpo, headers=cab, metodo="PATCH")
        accion = "regenerada — OJO: se perdieron las marcas anteriores"
    else:
        cuerpo, cab = _multipart({"name": nombre, "mimeType": HOJA,
                                  "parents": [carpeta_id]}, tmp, XLSX)
        sid = api("https://www.googleapis.com/upload/drive/v3/files"
                  "?uploadType=multipart&supportsAllDrives=true",
                  data=cuerpo, headers=cab)["id"]
        accion = "creada"
    print(f"✓ Hoja de estado {accion}: {nombre}")
    print(f"  https://docs.google.com/spreadsheets/d/{sid}/edit")
    print(sid)


if __name__ == "__main__":
    main()
